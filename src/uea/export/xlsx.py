"""XLSX writer for the tables of `tables.py` (openpyxl, MIT).

Every table is a sheet: title and state in the first rows, a header row with a filter, the data,
a sum row outside the filter and the notes. Values are written as numbers, not as formulas: the
code has already calculated them.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from uea import __version__
from uea.export.tables import Table

HEAD = 4
"""Row of the column headers."""


def write_xlsx(tables: list[Table], path: Path, project: str, batch: int | None) -> None:
    wb = Workbook()
    wb.remove(wb["Sheet"])
    state = f"Stand Batch {batch}" if batch is not None else "Stand ohne Verlauf"
    stamp = (
        f"{project} · {state} · aus dem UEA-Modell abgeleitet (UEA {__version__}),"
        " nicht geprüft und nicht unterzeichnet"
    )
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8E8E0")
    thin = Side(style="thin", color="888888")
    for t in tables:
        ws = wb.create_sheet(t.name)
        ws["A1"] = t.title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = stamp
        ws["A2"].font = Font(italic=True, color="666666")
        for i, c in enumerate(t.cols, 1):
            cell = ws.cell(HEAD, i, c.head)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = c.width
        for r, row in enumerate(t.rows, HEAD + 1):
            for i, (c, v) in enumerate(zip(t.cols, row, strict=True), 1):
                cell = ws.cell(r, i, v)
                if c.fmt is not None and v is not None:
                    cell.number_format = c.fmt
        last = HEAD + len(t.rows)
        ws.auto_filter.ref = f"A{HEAD}:{get_column_letter(len(t.cols))}{last}"
        ws.freeze_panes = ws.cell(HEAD + 1, 2)
        nxt = last + 1
        if t.total is not None:
            nxt = last + 2
            for i, (c, v) in enumerate(zip(t.cols, t.total, strict=True), 1):
                cell = ws.cell(nxt, i, v)
                cell.font = head_font
                cell.border = Border(top=thin)
                if c.fmt is not None and v is not None:
                    cell.number_format = c.fmt
        for k, note in enumerate(t.notes, nxt + 2):
            ws.cell(k, 1, note).font = Font(italic=True, color="666666")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{HEAD}:{HEAD}"
    wb.properties.creator = f"UEA {__version__}"
    wb.properties.title = f"{project}: Tabellen"
    wb.save(path)
