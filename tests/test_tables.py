"""Tables and the XLSX export, against hand-computed quantities of a 10 x 8 box house."""

import math
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.conftest import BOX, derive_text
from tests.test_geometry import ROOF, SLABS
from uea.core.values import Layers
from uea.derive import Derived
from uea.export.tables import Table, layers_de, tables, wall_qty
from uea.export.xlsx import HEAD, write_xlsx
from uea.packs.arch.geometry import ArchGeo, WallGeo

# EG z=0 fb=0.15, OG z=2.9: the slab above has its core top at 2.9-0.15 = 2.75 and its core
# underside at 2.55; the walls stand on -0.15, so they are 2.7 high.
H = 2.7
HOUSE = (
    BOX
    + SLABS
    + "door d1 w1 1x2.01 x=W+2 into=r1 hand=l\n"
    + "win f1 w3 1.2x1.2 x=W+3\n"
    + "niche ni1 w2.w 1x2.1 y=w1+2 d=0.1\n"
    + "room r1 EG living at=W+1,S+1 floor=FB\n"
)


def sheet(ts: list[Table], name: str) -> Table:
    return next(t for t in ts if t.name == name)


def rows_by_id(t: Table) -> dict[Any, list[Any]]:
    return {r[0]: r for r in t.rows}


def col(t: Table, head: str) -> int:
    return [c.head for c in t.cols].index(head)


def house(extra: str = "") -> Derived:
    d, rep = derive_text(HOUSE + extra)
    assert [i for i in rep.issues if i.code.startswith("E-")] == []
    return d


def test_walls_by_hand() -> None:
    walls = sheet(tables(house(), "EG"), "Wände")
    rows = rows_by_id(walls)
    length, gross = col(walls, "Länge m"), col(walls, "Fläche brutto m²")
    holes, net = col(walls, "Öffnungen m²"), col(walls, "Fläche netto m²")
    vol = col(walls, "Kernvolumen m³")
    # w1, w3: 10 m between the ends, w2, w4: 8 - 2 x 0.3 = 7.4 m between the faces of w1 and w3
    assert [rows[w][length] for w in ("w1", "w2", "w3", "w4")] == [10, 7.4, 10, 7.4]
    assert rows["w1"][gross] == pytest.approx(10 * H)
    assert rows["w2"][gross] == pytest.approx(7.4 * H)
    # the door 1 x 2.01 is in w1, the window 1.2 x 1.2 in w3; a niche is no hole
    assert rows["w1"][holes] == pytest.approx(2.01)
    assert rows["w3"][holes] == pytest.approx(1.44)
    assert rows["w2"][holes] == 0
    assert rows["w1"][net] == pytest.approx(10 * H - 2.01)
    assert rows["w1"][vol] == pytest.approx((10 * H - 2.01) * 0.3)
    assert rows["w1"][col(walls, "Höhe min m")] == pytest.approx(H)
    assert rows["w1"][col(walls, "Gesamtdicke m")] == pytest.approx(0.33)
    assert rows["w1"][col(walls, "Lage")] == "Außenwand"
    assert rows["w1"][col(walls, "Aufbau (cm, * Kern)")] == "gips 1 | *mw 30 | putz 2"
    assert walls.total is not None
    assert walls.total[length] == pytest.approx(34.8)
    assert walls.total[gross] == pytest.approx((20 + 14.8) * H)
    assert walls.total[holes] == pytest.approx(3.45)
    assert walls.total[net] == pytest.approx((20 + 14.8) * H - 3.45)
    assert walls.total[vol] == pytest.approx(((20 + 14.8) * H - 3.45) * 0.3)


def test_a_gable_wall_is_measured_under_its_top() -> None:
    # 10 m long, 2 m high at both ends and 4 m high in the middle: a 10 x 2 box and a triangle
    w = WallGeo(
        "w1", "EG", "h", 0.0, 0.3, 0.0, 10.0, Layers.parse("*mw:0.3"), False, "new", True,
        bottom=1.0, top=[(0.0, 3.0), (5.0, 5.0), (10.0, 3.0)],
    )  # fmt: skip
    q = wall_qty(ArchGeo(), w)
    assert q.gross == pytest.approx(10 * 2 + 0.5 * 10 * 2)
    assert (q.h_min, q.h_max) == (2.0, 4.0)
    assert q.volume == pytest.approx(30 * 0.3)


def test_a_wall_without_a_top_has_no_area_and_says_so() -> None:
    # nothing is above the OG: no slab, no roof
    d, _ = derive_text(BOX + "wall w5 OG AW y=S+ x=W..E\n")
    walls = sheet(tables(d, "OG"), "Wände")
    assert rows_by_id(walls)["w5"][col(walls, "Fläche brutto m²")] is None
    assert walls.notes[-1] == "Ohne Oberkante (keine Decke oder Dach darüber), ohne Fläche: w5"
    q = sheet(tables(d, "OG"), "Mengen")
    assert q.notes[-1] == "Wände ohne Oberkante gehen ohne Fläche ein: w5"


def test_openings() -> None:
    ts = tables(house(), "EG")
    ops = sheet(ts, "Öffnungen")
    rows = rows_by_id(ops)
    c = {h: col(ops, h) for h in ("Art", "Typ", "Breite m", "Brüstung m", "Sturz m", "Fläche m²")}
    assert (rows["d1"][c["Art"]], rows["d1"][c["Typ"]]) == ("Tür", "TI")
    assert rows["d1"][c["Fläche m²"]] == pytest.approx(2.01)
    assert rows["d1"][col(ops, "Anschlag")] == "links"
    assert rows["d1"][col(ops, "öffnet in")] == "r1"
    assert (rows["d1"][col(ops, "Raum 1")], rows["d1"][col(ops, "Raum 2")]) == ("r1", "außen")
    # a window hangs from the head height 2.26
    assert (rows["f1"][c["Art"]], rows["f1"][c["Typ"]]) == ("Fenster", "FE")
    assert rows["f1"][c["Brüstung m"]] == pytest.approx(2.26 - 1.2)
    assert rows["f1"][c["Sturz m"]] == pytest.approx(2.26)
    # a niche has no type, and its depth, and does not count in the sum
    assert rows["ni1"][c["Art"]] == "Nische"
    assert rows["ni1"][c["Typ"]] is None
    assert rows["ni1"][col(ops, "Tiefe m")] == pytest.approx(0.1)
    assert ops.total is not None and ops.total[c["Fläche m²"]] == pytest.approx(2.01 + 1.44)


def test_rooms() -> None:
    rooms = sheet(tables(house()), "Räume")
    (row,) = rooms.rows
    c = {h: col(rooms, h) for h in ("Nutzung", "Rohbaufläche m²", "Fertigfläche m²", "Bodenaufbau")}
    assert row[c["Nutzung"]] == "Wohnen"
    assert row[c["Rohbaufläche m²"]] == pytest.approx(9.4 * 7.4)
    assert row[c["Fertigfläche m²"]] == pytest.approx(9.38 * 7.38)
    assert row[c["Bodenaufbau"]] == "FB"
    assert rooms.total is not None
    assert rooms.total[c["Fertigfläche m²"]] == pytest.approx(9.38 * 7.38)


def test_slabs() -> None:
    slabs = sheet(tables(house()), "Decken")
    rows = rows_by_id(slabs)
    # the outline is the storey's outer faces, 10 x 8 less the plaster on the outside
    area = col(slabs, "Fläche netto m²")
    assert rows["sl1"][area] == rows["sl2"][area]
    assert rows["sl2"][col(slabs, "OK Rohdecke m")] == pytest.approx(2.75)
    assert rows["sl2"][col(slabs, "UK Rohdecke m")] == pytest.approx(2.55)
    assert rows["sl1"][col(slabs, "Kernvolumen m³")] == pytest.approx(rows["sl1"][area] * 0.2)


def test_a_gable_roof_is_measured_on_its_slope() -> None:
    d, _ = derive_text(ROOF)
    roofs = sheet(tables(d), "Dächer")
    (row,) = roofs.rows
    over = d.arch.roofs["rf1"].over.area
    assert row[col(roofs, "Grundfläche m²")] == pytest.approx(over)
    # both planes slope 45 degrees
    assert row[col(roofs, "Dachfläche m²")] == pytest.approx(over / math.cos(math.radians(45)))
    assert row[col(roofs, "Form")] == "Satteldach"
    assert row[col(roofs, "Neigung °")] == 45


def test_quantities_agree_with_the_tables() -> None:
    ts = tables(house())
    q = sheet(ts, "Mengen")
    walls = sheet(ts, "Wände")
    ops = sheet(ts, "Öffnungen")
    group, area = col(q, "Gruppe"), col(q, "Fläche m²")
    by_group: dict[str, float] = {}
    for r in q.rows:
        g, a = str(r[group]), r[area]
        by_group[g] = by_group.get(g, 0.0) + (a if isinstance(a, float) else 0.0)
    assert walls.total is not None and ops.total is not None
    assert by_group["Wand"] == pytest.approx(walls.total[col(walls, "Fläche netto m²")])
    assert by_group["Tür"] + by_group["Fenster"] == pytest.approx(ops.total[col(ops, "Fläche m²")])
    wall = next(r for r in q.rows if r[group] == "Wand")
    assert wall[col(q, "Typ")] == "AW"
    assert wall[col(q, "Anzahl")] == 4
    assert wall[col(q, "Länge m")] == pytest.approx(34.8)
    floor = next(r for r in q.rows if r[group] == "Boden")
    assert floor[area] == pytest.approx(9.38 * 7.38)


def test_status_is_kept_apart() -> None:
    d = house("wall w5 EG IW x=w4+4 y=w1..w3 demolish\n")
    ts = tables(d)
    walls = sheet(ts, "Wände")
    assert rows_by_id(walls)["w5"][col(walls, "Status")] == "Abbruch"
    assert walls.total is not None
    assert walls.total[col(walls, "Länge m")] == pytest.approx(34.8)  # w5 is not in the sum
    q = sheet(ts, "Mengen")
    demolished = [r for r in q.rows if r[col(q, "Status")] == "Abbruch"]
    assert [(r[col(q, "Gruppe")], r[col(q, "Typ")], r[col(q, "Anzahl")]) for r in demolished] == [
        ("Wand", "IW", 1)
    ]


def test_scope_is_one_storey() -> None:
    d = house("wall w5 OG AW on=w1\n")
    walls = sheet(tables(d, "OG"), "Wände")
    assert [r[0] for r in walls.rows] == ["w5"]
    assert walls.title == "Wände OG – Box"
    assert {t.name for t in tables(d, "OG")} <= {"Räume", "Wände", "Öffnungen", "Decken", "Mengen"}


def test_layers_in_cm_with_a_comma() -> None:
    assert layers_de(Layers.parse("putz:0.015,*ziegel:0.365,putz:0.02")) == (
        "putz 1,5 | *ziegel 36,5 | putz 2"
    )


def test_xlsx_file(tmp_path: Path) -> None:
    d = house()
    path = tmp_path / "h.xlsx"
    write_xlsx(tables(d), path, "Haus Box", 7)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Räume", "Wände", "Öffnungen", "Decken", "Mengen"]
    ws = wb["Wände"]
    assert ws["A1"].value == "Wände – Box"
    assert "Haus Box" in ws["A2"].value and "Stand Batch 7" in ws["A2"].value
    assert "nicht unterzeichnet" in ws["A2"].value
    heads = [c.value for c in ws[HEAD]]
    assert heads[:3] == ["Nr", "Geschoss", "Typ"]
    first = [c.value for c in ws[HEAD + 1]]
    assert first[0] == "w1" and first[7] == 10
    n = len(sheet(tables(d), "Wände").rows)
    assert ws.auto_filter.ref == f"A{HEAD}:P{HEAD + n}"
    assert ws.freeze_panes == f"B{HEAD + 1}"
    assert ws[f"A{HEAD + n + 2}"].value == "Summe ohne Abbruch"
    area = ws.cell(HEAD + 1, heads.index("Fläche brutto m²") + 1)
    assert area.number_format == "0.00"
